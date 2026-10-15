# -*- coding: utf-8 -*-
"""
Professional Excel Generator:
Title: ALB SSL Certificate Domain Missing in Route 53 Investigation Guide
Theme: Monotone Elegance (Charcoal, Slate, Mist Grey, Pure White)
Font: 100% Meiryo UI on all cells (including merged, blank, and headers)
Features: Freeze Panes, Auto-wrap, Custom row heights, Gridlines enabled.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_workbook(filename="ALB_SSL_Route53_Investigation_Guide.xlsx"):
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Palette Definition (Refined Monotone)
    # -------------------------------------------------------------
    HEX_TITLE_BG     = "1A1D20"   # Almost Black Carbon
    HEX_HEADER_BG    = "2D3136"   # Deep Charcoal
    HEX_SUBHEADER_BG = "495057"   # Slate Grey
    HEX_SECTION_BG   = "6C757D"   # Neutral Steel Grey
    HEX_ZEBRA        = "F8F9FA"   # Soft Off-White
    HEX_CARD         = "E9ECEF"   # Cool Pale Grey
    HEX_WHITE        = "FFFFFF"   # Crisp White
    HEX_BORDER_THIN  = "CED4DA"   # Elegant Light Border
    HEX_BORDER_MID   = "6C757D"   # Medium Border
    HEX_BORDER_DARK  = "212529"   # Dark Boundary Border
    HEX_TEXT_MAIN    = "212529"   # Body Text Charcoal
    HEX_TEXT_MUTED   = "495057"   # Subdued Text
    HEX_TEXT_WHITE   = "FFFFFF"   # White Text

    # Fonts
    font_title    = Font(name="Meiryo UI", size=14, bold=True, color=HEX_TEXT_WHITE)
    font_subtitle = Font(name="Meiryo UI", size=9.5, italic=True, color="E0E0E0")
    font_section  = Font(name="Meiryo UI", size=10.5, bold=True, color=HEX_TEXT_WHITE)
    font_header   = Font(name="Meiryo UI", size=9.5, bold=True, color=HEX_TEXT_WHITE)
    font_sub_hdr  = Font(name="Meiryo UI", size=9, bold=True, color=HEX_TEXT_WHITE)
    font_body     = Font(name="Meiryo UI", size=9, color=HEX_TEXT_MAIN)
    font_bold     = Font(name="Meiryo UI", size=9, bold=True, color=HEX_TEXT_MAIN)
    font_code     = Font(name="Meiryo UI", size=8.5, color="1B1E1C")
    font_note     = Font(name="Meiryo UI", size=8, italic=True, color=HEX_TEXT_MUTED)

    # Fills
    fill_title    = PatternFill(start_color=HEX_TITLE_BG, end_color=HEX_TITLE_BG, fill_type="solid")
    fill_header   = PatternFill(start_color=HEX_HEADER_BG, end_color=HEX_HEADER_BG, fill_type="solid")
    fill_sub_hdr  = PatternFill(start_color=HEX_SUBHEADER_BG, end_color=HEX_SUBHEADER_BG, fill_type="solid")
    fill_section  = PatternFill(start_color=HEX_SECTION_BG, end_color=HEX_SECTION_BG, fill_type="solid")
    fill_zebra    = PatternFill(start_color=HEX_ZEBRA, end_color=HEX_ZEBRA, fill_type="solid")
    fill_card     = PatternFill(start_color=HEX_CARD, end_color=HEX_CARD, fill_type="solid")
    fill_white    = PatternFill(start_color=HEX_WHITE, end_color=HEX_WHITE, fill_type="solid")

    # Borders
    thin_s  = Side(border_style="thin", color=HEX_BORDER_THIN)
    mid_s   = Side(border_style="medium", color=HEX_BORDER_MID)
    dark_d  = Side(border_style="double", color=HEX_BORDER_DARK)

    border_cell   = Border(left=thin_s, right=thin_s, top=thin_s, bottom=thin_s)
    border_header = Border(left=thin_s, right=thin_s, top=thin_s, bottom=mid_s)

    # Alignments
    align_center   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left     = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_top_mid  = Alignment(horizontal="center", vertical="top", wrap_text=True)

    def apply_range_style(ws, min_col, min_row, max_col, max_row, font=None, fill=None, border=None, alignment=None):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                if font: cell.font = font
                if fill: cell.fill = fill
                if border: cell.border = border
                if alignment: cell.alignment = alignment

    # =============================================================
    # SHEET 1: 概要と小学生向け解説
    # =============================================================
    ws1 = wb.active
    ws1.title = "概要と小学生向け解説"
    ws1.views.sheetView[0].showGridLines = True
    
    # Title
    ws1.merge_cells("A1:G2")
    ws1["A1"] = "ALBのSSL証明書ドメインがRoute 53に存在しない謎：総合分析レポート"
    apply_range_style(ws1, 1, 1, 7, 2, font=font_title, fill=fill_title, alignment=align_center)
    ws1.row_dimensions[1].height = 20
    ws1.row_dimensions[2].height = 20

    ws1.merge_cells("A3:G3")
    ws1["A3"] = "〜 クラウドにおけるDNS（案内看板）・証明書（身分証）・ロードバランサー（受付）の独立構造と10大可能性 〜"
    apply_range_style(ws1, 1, 3, 7, 3, font=font_subtitle, fill=fill_sub_hdr, alignment=align_center)
    ws1.row_dimensions[3].height = 22

    # Section 1: 小学生にもわかるたとえ話
    r = 5
    ws1.merge_cells(f"A{r}:G{r}")
    ws1[f"A{r}"] = "◆ 【小学生でも1分でわかる！】案内所のお姉さんと街の電話帳のナゾ"
    apply_range_style(ws1, 1, r, 7, r, font=font_section, fill=fill_section, alignment=align_left)
    ws1.row_dimensions[r].height = 26
    r += 1

    analogy_data = [
        ("ITの世界の用語", "小学生向け身近なたとえ", "役割とたとえの解説"),
        ("ALB (ロードバランサー)", "巨大デパートの『総合案内所のお姉さん』", "やってきた大勢のお客さんを見て、『おもちゃ売り場は3階へ』『レストランは最上階へ』とテキパキ振り分けてくれる係。"),
        ("SSL/TLS 証明書 (ACM)", "首から下げた『キラキラの公式名札』", "『私は正式な「ひまわりデパート」のスタッフです！偽物ではありません』と証明する安心の身分証。これがないと怪しまれてお客さんが逃げてしまう。"),
        ("Route 53 (DNSサービス)", "街の公式『住所録・電話帳（104番）』", "『ひまわりデパートに行きたい人は、駅前通り5丁目を右に曲がってね』と場所（IPアドレス）を教えてくれる案内板。"),
        ("今回の事件の状況", "「名札はあるのに、公式住所録に載っていない！」", "お姉さんは『ひまわりデパート』の名札を誇らしげにつけて立っているのに、街の公式住所録（Route 53）をいくら調べても『ひまわりデパート』の住所が書いていない！？という不思議な状態です。")
    ]
    for idx, item in enumerate(analogy_data):
        ws1.row_dimensions[r].height = 32
        is_hdr = (idx == 0)
        ws1[f"A{r}"] = item[0]
        ws1[f"B{r}"] = item[1]
        ws1.merge_cells(f"C{r}:G{r}")
        ws1[f"C{r}"] = item[2]
        
        if is_hdr:
            apply_range_style(ws1, 1, r, 2, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
            apply_range_style(ws1, 3, r, 7, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
        else:
            apply_range_style(ws1, 1, r, 1, r, font=font_bold, fill=fill_card, border=border_cell, alignment=align_center)
            apply_range_style(ws1, 2, r, 2, r, font=font_bold, fill=fill_white, border=border_cell, alignment=align_left)
            apply_range_style(ws1, 3, r, 7, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_top_left)
        r += 1

    r += 1
    # Section 2: なぜそんなことが起きるの？
    ws1.merge_cells(f"A{r}:G{r}")
    ws1[f"A{r}"] = "◆ 【なぜそんなことが起きるの？】5つの身近なたとえで納得！"
    apply_range_style(ws1, 1, r, 7, r, font=font_section, fill=fill_section, alignment=align_left)
    ws1.row_dimensions[r].height = 26
    r += 1

    reasons_kid = [
        ("たとえ話のパターン", "ITの世界の本当の姿", "わかりやすい説明"),
        ("① 別の会社の電話帳に載っている", "外部DNS（Cloudflare / お名前.com等）", "街の公式電話帳（Route 53）ではなく、海外の電話帳や民間会社の電話帳に住所が載っているから。デパートにはちゃんと行けます。"),
        ("② 本社の秘密の電話帳に載っている", "マルチアカウント構成（統括アカウントのDNS）", "支店のお店には電話帳が置いておらず、本社の大きな部屋にある共通電話帳で一括管理されているから。"),
        ("③ 駅前の無料シャトルバスに乗る", "CDN前段配置（CloudFront / WAF）", "お客さんはデパートに直接歩いていかず、駅前の高速シャトルバスに乗って行く。看板はバス乗り場に向いている。"),
        ("④ 明日オープンなので看板は夜につける", "構築中・移行中（DNS切り替え前）", "お店もお姉さんの名札も準備万端だけど、グランドオープンの明日までは電話帳に載せない計画だから。"),
        ("⑤ 昔の名札をはずし忘れている", "消し忘れ（ゾンビ設定・ゴースト証明書）", "以前のお店はもう閉じたのに、案内所のお姉さんが昔の名札を名札立てに差しっぱなしにしていたから。")
    ]
    for idx, item in enumerate(reasons_kid):
        ws1.row_dimensions[r].height = 28
        is_hdr = (idx == 0)
        ws1[f"A{r}"] = item[0]
        ws1[f"B{r}"] = item[1]
        ws1.merge_cells(f"C{r}:G{r}")
        ws1[f"C{r}"] = item[2]
        
        if is_hdr:
            apply_range_style(ws1, 1, r, 2, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
            apply_range_style(ws1, 3, r, 7, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
        else:
            apply_range_style(ws1, 1, r, 1, r, font=font_bold, fill=fill_card, border=border_cell, alignment=align_left)
            apply_range_style(ws1, 2, r, 2, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_left)
            apply_range_style(ws1, 3, r, 7, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_top_left)
        r += 1

    r += 1
    # Section 3: エンジニア向け結論サマリ
    ws1.merge_cells(f"A{r}:G{r}")
    ws1[f"A{r}"] = "◆ 【アーキテクチャ総括】クラウドにおける独立性と正常性判定"
    apply_range_style(ws1, 1, r, 7, r, font=font_section, fill=fill_section, alignment=align_left)
    ws1.row_dimensions[r].height = 26
    r += 1

    summary_bullets = [
        "1. コントロールプレーンの完全分離: AWSにおいて証明書（ACM）の発行・ALBリスナーへのアタッチと、DNSレコード（Route 53）の作成は完全に独立したAPIと権限境界を持つ。ALBは証明書のドメインがRoute 53に引けるかを一切関知しない。",
        "2. ACM DNS検証の仕組み: ACMの検証は世界中どの権威DNS（Route 53、Cloudflare、オンプレミスBIND等）に検証用CNAMEを追加しても成立する。Route 53を使っていなくても正規の証明書が発行されALBに正常にバインドされる。",
        "3. エンタープライズ設計の最頻出パターン: 実務において『ALBに証明書があるのにRoute 53にない』事象の8割以上は、①『Cloudflare等の外部DNS利用』、②『AWS OrganizationsによるNetworkアカウント集約』、③『CloudFrontのOrigin暗号化用』のいずれかであり、極めて正常な設計である。"
    ]
    for sb in summary_bullets:
        ws1.row_dimensions[r].height = 34
        ws1.merge_cells(f"A{r}:G{r}")
        ws1[f"A{r}"] = sb
        apply_range_style(ws1, 1, r, 7, r, font=font_bold, fill=fill_card, border=border_cell, alignment=align_top_left)
        r += 1

    col_widths_ws1 = [24, 26, 18, 18, 18, 18, 18]
    for idx, w in enumerate(col_widths_ws1, start=1):
        ws1.column_dimensions[get_column_letter(idx)].width = w

    # =============================================================
    # SHEET 2: 10大可能性一覧マトリクス
    # =============================================================
    ws2 = wb.create_sheet(title="10大可能性一覧マトリクス")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "A3"  # Freeze headers
    
    ws2.merge_cells("A1:I1")
    ws2["A1"] = "ALBのSSL証明書ドメインがRoute 53に存在しない全シナリオ比較マトリクス"
    apply_range_style(ws2, 1, 1, 9, 1, font=font_title, fill=fill_title, alignment=align_center)
    ws2.row_dimensions[1].height = 36

    headers_s2 = [
        ("No", 6),
        ("分類カテゴリ", 16),
        ("シナリオ・原因の仮説", 28),
        ("発生頻度", 10),
        ("なぜALBに証明書があるのか？", 32),
        ("なぜRoute 53にレコードがないのか？", 32),
        ("正常/異常", 12),
        ("実務での典型例・背景", 30),
        ("確認・切り分けコマンド", 28)
    ]
    
    ws2.row_dimensions[2].height = 26
    for col_idx, (h_name, width) in enumerate(headers_s2, start=1):
        cell = ws2.cell(row=2, column=col_idx, value=h_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
        ws2.column_dimensions[get_column_letter(col_idx)].width = width

    scenarios_data = [
        (
            1, "外部DNS運用", "他社DNS（Cloudflare / お名前.com等）を利用", "極めて高",
            "外部DNSにACM検証用CNAMEを手動登録して証明書を発行し、ALBにアタッチした。",
            "ドメインの権威DNSがRoute 53ではなく外部DNSにあり、外部側でALBのDNS名へCNAMEしている。",
            "正常（標準的）", "ドメイン管理を会社規定で特定レジストラやCloudflareに集約している企業で多発。",
            "whois example.com\nnslookup -type=ns example.com"
        ),
        (
            2, "マルチアカウント", "組織共通DNSアカウント（Network / Shared）に集約", "極めて高",
            "ALBが所在するサービス開発用AWSアカウント内でACM証明書を発行・バインドした。",
            "Route 53パブリックホストゾーンは組織共通ネットワークアカウントにあり、作業アカウントにない。",
            "正常（統制設計）", "AWS Control TowerやOrganizationsを導入したエンタープライズの標準構成。",
            "AWS Organizations構成確認\n別アカウントのRoute 53調査"
        ),
        (
            3, "CDN前段配置", "CloudFront + ALB（End-to-End暗号化のオリジン）", "極めて高",
            "CloudFrontからALB間の通信をHTTPSで安全に保護するため、ALB側にも同一ドメイン証明書が必要。",
            "Route 53のエイリアスレコードは「CloudFront」を向いており、ALBへ直接向けるレコードはない。",
            "正常（推奨設計）", "静的キャッシュ配信、エッジWAF適用、DDoS対策でCloudFrontを前段に置くWebシステム。",
            "Route 53でCloudFrontエイリアス探索\nCloudFrontのOrigin設定確認"
        ),
        (
            4, "外部WAF/CDN", "SaaS型WAF（Akamai / Imperva / Fastly等）経由", "中〜高",
            "WAFプロキシからALB間のTLS通信を暗号化・SNI検証するためにALBリスナーに正規証明書が必要。",
            "パブリックDNSはWAFベンダーのエッジIP/CNAMEを向いており、ALBへの直接レコードは非公開。",
            "正常（企業標準）", "全社統一セキュリティ基盤としてAkamai等のCDN/WAFを一元導入している金融・大手企業。",
            "dig +trace example.com\n返却IPのWHOIS/ASNを確認"
        ),
        (
            5, "閉域・社内DNS", "プライベートホストゾーン（PHZ）または社内AD DNS", "中",
            "社内イントラやVPC間通信でもブラウザのSSL警告を出さないため、正規証明書をALBにアタッチ。",
            "パブリックには一切レコードを作らず、特定VPC紐付けPHZまたはDirect Connect経由の社内DNSにある。",
            "正常（閉域網）", "社内業務システム、マイクロサービスのバックエンドAPI、管理画面など。",
            "aws route53 list-hosted-zones\n社内AD DNSの正引き確認"
        ),
        (
            6, "証明書仕様", "ワイルドカード (*.example.com) やマルチドメイン (SAN)", "高",
            "1枚の証明書で将来追加予定のサブドメインや他アプリもカバーするためALBに適用。",
            "Route 53には現在稼働中のサブドメインしか作っておらず、未稼働のサブドメインレコードはない。",
            "正常（仕様通り）", "複数システムを1つのワイルドカード証明書で効率よく運用しているケース。",
            "openssl s_client -connect <ALB-DNS>:443\n-servername <Domain> でSAN一覧確認"
        ),
        (
            7, "構築フェーズ", "リリース前・移行中・Blue/Greenデプロイ（切替前）", "中",
            "インフラ構築（Terraform/CloudFormation）でALB・ECS・証明書を先行構築し事前テスト中。",
            "本番DNSの切り替え（DNSカットオーバー）が未実施のため、Route 53レコードは旧環境のまま。",
            "正常（過渡期）", "新規サービスのローンチ準備中、またはオンプレミスからAWSへの移行作業中。",
            "リリース計画書・切替手順書確認\nhostsファイルを書き換えて疎通テスト"
        ),
        (
            8, "構成不備・残骸", "サービス廃止後の消し忘れ（ゾンビ設定・ゴースト証明書）", "中",
            "以前そのドメインで運用していたが、サービス終了時にRoute 53レコードだけ削除しALB側を放置。",
            "DNSレコードは削除されたが、ALBリスナーのSNI追加証明書からデタッチする作業が漏れていた。",
            "要対応（残骸）", "頻繁にサブドメインの追加・削除が発生する開発環境や、退職者が残した設定。",
            "CloudWatch ALB RequestCount確認\nアクセス0ならリスナーから削除"
        ),
        (
            9, "ALB仕様制約", "デフォルト証明書（フォールバック用）として残存", "中〜低",
            "ALB HTTPSリスナーは仕様上、最低1つの「デフォルト証明書」が必須。作成時に初期選択した。",
            "実際の運用はSNIで後から追加した別証明書のドメインで行われており、デフォルト証明書は不使用。",
            "正常/整理推奨", "当初は別ドメインでALBを作り、後から本番ドメインをSNI追加して運用しているケース。",
            "ALBマネジメントコンソールで\n「デフォルト証明書」と「SNI証明書」を比較"
        ),
        (
            10, "ヒューマンエラー", "タイポ（ドメインのスペルミス）または誤ったホストゾーン", "低〜中",
            "ACM証明書をリクエストする際、スペルミスしたドメイン（例: .co.jp と .comの間違い）で発行した。",
            "正しいドメインのRoute 53ホストゾーンを探しているが、そもそも証明書側の文字列が間違っている。",
            "異常（設定ミス）", "手動コンソール作成時や、パラメータファイルの定義間違いによるもの。",
            "証明書のドメイン名と探しているドメイン\nの文字コード・完全一致比較"
        )
    ]

    for row_idx, sc in enumerate(scenarios_data, start=3):
        ws2.row_dimensions[row_idx].height = 42
        fill_curr = fill_zebra if row_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(sc, start=1):
            cell = ws2.cell(row=row_idx, column=c_idx, value=val)
            cell.font = font_body
            cell.border = border_cell
            cell.fill = fill_curr
            if c_idx == 1:
                cell.alignment = align_center
                cell.font = font_bold
            elif c_idx in (2, 4, 7):
                cell.alignment = align_center
                if c_idx == 7 and "要対応" in str(val):
                    cell.font = font_bold
            else:
                cell.alignment = align_top_left

    # =============================================================
    # SHEET 3: 動作原理とメカニズム深堀り
    # =============================================================
    ws3 = wb.create_sheet(title="動作原理とメカニズム深堀り")
    ws3.views.sheetView[0].showGridLines = True
    ws3.freeze_panes = "A3"
    
    ws3.merge_cells("A1:G1")
    ws3["A1"] = "DNS名前解決 × TLSハンドシェイク × ALBアーキテクチャの完全解剖"
    apply_range_style(ws3, 1, 1, 7, 1, font=font_title, fill=fill_title, alignment=align_center)
    ws3.row_dimensions[1].height = 36

    r = 3
    # Step 1
    ws3.merge_cells(f"A{r}:G{r}")
    ws3[f"A{r}"] = "◆ ① コントロールプレーンの独立性：なぜRoute 53なしで証明書が貼れるのか？"
    apply_range_style(ws3, 1, r, 7, r, font=font_section, fill=fill_section, alignment=align_left)
    ws3.row_dimensions[r].height = 26
    r += 1

    tech_points = [
        ("フェーズ", "AWSサービス", "処理内容", "Route 53への依存性", "詳細解説と動作原理"),
        ("1. 証明書リクエスト", "AWS Certificate Manager (ACM)", "ドメイン名（例: api.example.com）を指定して発行申請", "全く不要", "ACMは申請を受け付けると、所有権検証用のランダムなCNAME文字列（_xxxx.api.example.com CNAME _yyyy.acm-validations.aws.）を生成する。この時点ではDNSに何も存在しなくてよい。"),
        ("2. ドメイン所有権検証", "外部DNSまたはRoute 53", "DNS検証（またはEメール検証）の実行", "Route 53である必要なし", "世界中のどの権威DNSであっても、指定された検証用CNAMEレコードが引ければACMは『所有者である』と認定し証明書を【発行済み（Issued）】にする。Route 53を一切使わなくても完了する。"),
        ("3. ALBへのバインド", "Application Load Balancer (ALB)", "HTTPSリスナーに証明書ARNを紐付け", "全く不要", "ALBのAWS API（ModifyListener / AddListenerCertificates）は、ACM証明書がIssued状態であれば無条件にリスナーへ設定可能。ドメインへのルーティングが実在するかは一切検査しない。"),
        ("4. クライアントからの通信", "ブラウザ・HTTPクライアント", "ドメインへHTTPSアクセスを試行", "DNSレコードが必須！", "クライアントが通信を開始するには『ドメイン名 → IPアドレス』の解決が必須。ここでDNSレコードが存在しないと『DNS_PROBE_FINISHED_NXDOMAIN』となり、ALBにパケットすら届かない。")
    ]

    for idx, tp in enumerate(tech_points):
        ws3.row_dimensions[r].height = 36
        is_h = (idx == 0)
        ws3[f"A{r}"] = tp[0]
        ws3[f"B{r}"] = tp[1]
        ws3[f"C{r}"] = tp[2]
        ws3[f"D{r}"] = tp[3]
        ws3.merge_cells(f"E{r}:G{r}")
        ws3[f"E{r}"] = tp[4]
        
        if is_h:
            apply_range_style(ws3, 1, r, 4, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
            apply_range_style(ws3, 5, r, 7, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
        else:
            apply_range_style(ws3, 1, r, 1, r, font=font_bold, fill=fill_card, border=border_cell, alignment=align_center)
            apply_range_style(ws3, 2, r, 2, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_left)
            apply_range_style(ws3, 3, r, 3, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_left)
            apply_range_style(ws3, 4, r, 4, r, font=font_bold, fill=fill_white, border=border_cell, alignment=align_center)
            apply_range_style(ws3, 5, r, 7, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_top_left)
        r += 1

    r += 1
    # Step 2
    ws3.merge_cells(f"A{r}:G{r}")
    ws3[f"A{r}"] = "◆ ② ALBにおけるTLSハンドシェイクとSNI（Server Name Indication）選択アルゴリズム"
    apply_range_style(ws3, 1, r, 7, r, font=font_section, fill=fill_section, alignment=align_left)
    ws3.row_dimensions[r].height = 26
    r += 1

    sni_steps = [
        ("Step", "通信主体", "プロトコル / パケット", "暗号化状態", "ALB内部の処理ロジックと動作イメージ"),
        ("1", "Client → ALB", "TCP 3-way Handshake (SYN / SYN-ACK / ACK)", "平文", "TCPポート443へのコネクション確立。この段階ではドメイン名もHTTPヘッダーもまだ一切送られていない。ALBは通信元IPのみを認識。"),
        ("2", "Client → ALB", "TLS ClientHello (SNI拡張ヘッダー含む)", "平文 (平文SNI)", "ブラウザが接続先ホスト名（例: api.example.com）を平文のServer Name Indication拡張フィールドに載せて送信。"),
        ("3", "ALB 内部処理", "証明書選択エンジン (SNI Evaluation)", "内部処理", "ALBは登録済み証明書（最大25個）から探索：\n① 完全一致 (api.example.com)\n② 最長一致ワイルドカード (*.example.com)\n③ 合致がなければ【デフォルト証明書】を選択。"),
        ("4", "ALB → Client", "TLS ServerHello + Certificate + Key Exchange", "暗号化開始", "選択したSSL証明書をクライアントへ返却。クライアントは認証局（CA）を検証し、暗号鍵（共通鍵）を安全に交換。"),
        ("5", "Client ↔ ALB", "暗号化HTTP通信 (GET / Host: api.example.com)", "完全暗号化", "TLSトンネルが確立された後、初めて暗号化されたHTTP Hostヘッダーが到達。ALBはリスナールールに基づいてターゲットグループへ転送する。")
    ]

    for idx, st in enumerate(sni_steps):
        ws3.row_dimensions[r].height = 36
        is_h = (idx == 0)
        ws3[f"A{r}"] = st[0]
        ws3[f"B{r}"] = st[1]
        ws3[f"C{r}"] = st[2]
        ws3[f"D{r}"] = st[3]
        ws3.merge_cells(f"E{r}:G{r}")
        ws3[f"E{r}"] = st[4]
        
        if is_h:
            apply_range_style(ws3, 1, r, 4, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
            apply_range_style(ws3, 5, r, 7, r, font=font_header, fill=fill_header, border=border_header, alignment=align_center)
        else:
            apply_range_style(ws3, 1, r, 1, r, font=font_bold, fill=fill_card, border=border_cell, alignment=align_center)
            apply_range_style(ws3, 2, r, 2, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_center)
            apply_range_style(ws3, 3, r, 3, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_left)
            apply_range_style(ws3, 4, r, 4, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_center)
            apply_range_style(ws3, 5, r, 7, r, font=font_body, fill=fill_white, border=border_cell, alignment=align_top_left)
        r += 1

    col_widths_ws3 = [8, 16, 28, 16, 22, 22, 22]
    for idx, w in enumerate(col_widths_ws3, start=1):
        ws3.column_dimensions[get_column_letter(idx)].width = w

    # =============================================================
    # SHEET 4: 歴史と背景・必要になった経緯
    # =============================================================
    ws4 = wb.create_sheet(title="歴史と背景・技術進化")
    ws4.views.sheetView[0].showGridLines = True
    ws4.freeze_panes = "A3"
    
    ws4.merge_cells("A1:F1")
    ws4["A1"] = "インターネットとクラウドの進化史：DNS・SSL/TLS・ロードバランサーが辿った必然の軌跡"
    apply_range_style(ws4, 1, 1, 6, 1, font=font_title, fill=fill_title, alignment=align_center)
    ws4.row_dimensions[1].height = 36

    headers_s4 = [
        ("年代 / 時期", 14),
        ("技術・プロトコル", 20),
        ("当時の課題・社会背景", 34),
        ("ブレイクスルーと発明された仕組み", 34),
        ("現代への影響・ALBとの関係", 28),
        ("技術カテゴリ", 12)
    ]
    
    ws4.row_dimensions[2].height = 26
    for c_idx, (h_name, w) in enumerate(headers_s4, start=1):
        cell = ws4.cell(row=2, column=c_idx, value=h_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
        ws4.column_dimensions[get_column_letter(c_idx)].width = w

    history_data = [
        (
            "1970年代〜1983年", "HOSTS.TXT から DNS へ",
            "ARPANET時代はスタンフォード研究所が単一のファイル『HOSTS.TXT』を手動更新し全員がFTPでダウンロード。接続機器激増で破綻寸前に。",
            "1983年 ポール・モカペトリスが『DNS』を考案。階層型（Tree構造）の自律分散データベースにより、各組織が自分のドメインを独自管理可能に。",
            "ドメイン名が世界規模で独立分散管理される大原則が確立。Route 53もこの世界分散DNSの1つの権威サーバーに過ぎない。", "DNS基盤"
        ),
        (
            "1994年〜1995年", "SSL 2.0 / 3.0 の誕生",
            "Netscape Navigatorの登場でWeb上でクレジットカード情報などをやり取りする商業利用（Eコマース）が急拡大。通信の盗聴と改ざんが深刻化。",
            "Netscape社がSSL（Secure Sockets Layer）を開発。暗号化、サーバー認証（証明書）、完全性検証をワンセットで提供。",
            "現代のHTTPS通信の基礎。ロードバランサー（ALB含む）が証明書を提示して信頼性を担保する構造が誕生。", "セキュリティ"
        ),
        (
            "1999年〜2003年", "TLS標準化と SNI の発明",
            "従来のSSL/TLSは『1つのIPアドレスに1つの証明書』しか返せなかった。バーチャルホストで多数のWebサイトを収容するとIPv4が即座に枯渇。",
            "RFC 3546 / RFC 6066で『SNI (Server Name Indication)』を定義。TLS接続開始時にアクセス先ドメイン名を送ることで、1つのIPで複数証明書の共存が可能に。",
            "ALBが単一のIP/リスナーで最大25個の異なるドメイン証明書を保持できるのは、このSNI技術の賜物。", "TLS拡張"
        ),
        (
            "2009年〜2010年", "AWS ELB (CLB) と Route 53",
            "クラウド草創期。EC2の急激なスケールアウトに伴い、DNSとロードバランサーの自動追随・高可用性が求められた。",
            "2009年にClassic Load Balancer (CLB)、2010年にRoute 53が登場。Route 53独自の『Aliasレコード』でELBの動的IPに追随可能に。",
            "AWS内でのDNSとロードバランサーの統合が完成。しかし当時はCLBの仕様上、複数ドメイン証明書の共存が極めて困難だった。", "AWS黎明期"
        ),
        (
            "2016年〜2017年", "ALB登場と SNI サポート",
            "マイクロサービスやDocker/ECSコンテナの台頭で、URLパスやホストヘッダーによるきめ細かなL7ルーティングとコスト集約が不可欠に。",
            "2016年にALBがリリース。2017年に待望の『SNIサポート』が追加され、1つのALBに最大25個の証明書を登録可能になった。",
            "『1つのALBに多種多様なドメイン証明書が集約される』環境が完成。今回の『別アカウントや他社DNSのドメインが相乗りする現象』の土台となった。", "ALB革命"
        ),
        (
            "2018年〜現在", "Zero Trust とマルチCDN時代",
            "境界防御モデルの崩壊。エッジ（CloudFront / Cloudflare等）でのWAF・DDoS対策が必須化し、オリジン間も含めたEnd-to-End暗号化が常識化。",
            "エッジとオリジン（ALB）の両方に正規証明書を配置する設計が定着。またマルチアカウント管理（Control Tower）でDNSとアプリが分離。",
            "『パブリックDNSは別事業者や別アカウントにあり、ALBにはオリジン暗号化用証明書だけが載っている』という本問の現代的アーキテクチャが確立。", "最新トレンド"
        )
    ]

    for row_idx, h in enumerate(history_data, start=3):
        ws4.row_dimensions[row_idx].height = 42
        fill_curr = fill_zebra if row_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(h, start=1):
            cell = ws4.cell(row=row_idx, column=c_idx, value=val)
            cell.font = font_body
            cell.border = border_cell
            cell.fill = fill_curr
            if c_idx in (1, 6):
                cell.alignment = align_center
                if c_idx == 1:
                    cell.font = font_bold
            else:
                cell.alignment = align_top_left

    # =============================================================
    # SHEET 5: 実践調査・トラブルシューティング手順書
    # =============================================================
    ws5 = wb.create_sheet(title="実践トラブルシューティング手順")
    ws5.views.sheetView[0].showGridLines = True
    ws5.freeze_panes = "A3"
    
    ws5.merge_cells("A1:G1")
    ws5["A1"] = "現場のエンジニアが即座に実行できる原因特定コマンドと調査フロー"
    apply_range_style(ws5, 1, 1, 7, 1, font=font_title, fill=fill_title, alignment=align_center)
    ws5.row_dimensions[1].height = 36

    headers_s5 = [
        ("Step", 8),
        ("調査項目", 18),
        ("実行コマンド / 調査手順", 42),
        ("出力結果の着眼点", 28),
        ("判定結果と次のアクション", 30),
        ("使用ツール", 12),
        ("優先度", 10)
    ]
    
    ws5.row_dimensions[2].height = 26
    for c_idx, (h_name, w) in enumerate(headers_s5, start=1):
        cell = ws5.cell(row=2, column=c_idx, value=h_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header
        ws5.column_dimensions[get_column_letter(c_idx)].width = w

    troubleshooting_steps = [
        (
            "Step 1", "権威DNSの特定",
            "dig NS <ドメイン名> +short\nnslookup -type=ns <ドメイン名>",
            "返却されるネームサーバーが「awsdns-xx」か「cloudflare.com」「gmo」「other」か？",
            "awsdns以外なら【外部DNS運用（シナリオ1）】確定。外部DNS管理者に設定を確認。",
            "dig / CLI", "最高"
        ),
        (
            "Step 2", "現在の名前解決確認",
            "dig A <ドメイン名> +trace\nnslookup <ドメイン名>",
            "そもそもIPアドレスが返るか？CNAMEでCloudFrontやALBにチェーンしているか？",
            "CloudFrontのDNS名なら【シナリオ3】。NXDOMAIN（レコードなし）なら【シナリオ7 or 8】。",
            "dig / CLI", "最高"
        ),
        (
            "Step 3", "Route 53全ゾーン検索",
            "aws route53 list-hosted-zones --query \"HostedZones[?contains(Name, 'example.com')]\"",
            "現在のAWSアカウント内にプライベートゾーンや別名ゾーンが存在するか？",
            "空配列（[]）なら、マルチアカウント環境で別アカウントに親ゾーンがあるか社内DNS。",
            "AWS CLI", "高"
        ),
        (
            "Step 4", "ALB実機証明書の検証",
            "openssl s_client -connect <ALB-DNS名>:443 -servername <ドメイン名> -showcerts",
            "Server certificate の CN / SAN に確かにそのドメインが含まれているか？",
            "含まれていればALB側のバインドは正常。ワイルドカードや複数ドメインも確認可能。",
            "OpenSSL", "高"
        ),
        (
            "Step 5", "ALBトラフィック確認",
            "CloudWatch Metrics:\nAWS/ApplicationELB -> RequestCount, TargetResponseTime",
            "対象ALBに対して直近7日間〜30日間でリクエストが実際に流入しているか？",
            "リクエストが0なら【シナリオ8：不要な残骸設定】の可能性濃厚。安全に削除検討。",
            "CloudWatch", "中"
        ),
        (
            "Step 6", "CloudFront / WAF確認",
            "aws cloudfront list-distributions --query \"DistributionList.Items[*].Aliases.Items\"",
            "自社CloudFrontディストリビューションのCNAME（代替ドメイン）に登録されているか？",
            "合致するものがあれば【シナリオ3：CloudFrontオリジン構成】。Route 53はCloudFrontを向くのが正解。",
            "AWS CLI", "中"
        )
    ]

    for row_idx, tb in enumerate(troubleshooting_steps, start=3):
        ws5.row_dimensions[row_idx].height = 42
        fill_curr = fill_zebra if row_idx % 2 == 0 else fill_white
        for c_idx, val in enumerate(tb, start=1):
            cell = ws5.cell(row=row_idx, column=c_idx, value=val)
            cell.font = font_body
            cell.border = border_cell
            cell.fill = fill_curr
            if c_idx in (1, 6, 7):
                cell.alignment = align_center
                if c_idx == 1:
                    cell.font = font_bold
            else:
                cell.alignment = align_top_left

    # -------------------------------------------------------------
    # Final Pass: Ensure 100% Meiryo UI on EVERY cell across ALL sheets
    # -------------------------------------------------------------
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if cell.font:
                    if cell.font.name != "Meiryo UI":
                        # preserve existing attributes, only replace font name
                        cell.font = Font(
                            name="Meiryo UI",
                            size=cell.font.size,
                            bold=cell.font.bold,
                            italic=cell.font.italic,
                            color=cell.font.color
                        )

    # Save workbook
    wb.save(filename)
    print(f"Workbook successfully saved: {filename}")

if __name__ == "__main__":
    generate_workbook()
